"""Bytes in, text segments out.

WHAT IS SUPPORTED IS DEFINED BY DocumentService's ALLOW-LIST, not by what happens to be
uploaded: .pdf .txt .md .markdown .csv .docx .xlsx, 10MB max. The corpus contains no DOCX
at all — building from the corpus would have left a hole that the first .docx upload fell
into, and it would have looked like a chunker bug rather than a missing extractor.

DISPATCH IS ON CONTENT, THEN EXTENSION — never on the recorded MIME type. Three documents
in the live database are stored as `application/octet-stream` and are actually a .md and
.xlsx. The stored filename has no extension either (files are saved as `uploads/<uuid>`),
which is why openpyxl refused them outright: it picks its parser from the extension and saw
none. Python receives the real filename over multipart for exactly this reason.

A SEGMENT IS THE UNIT THAT CARRIES METADATA. A PDF page, a Markdown heading section, a
spreadsheet sheet. Chunks are packed WITHIN a segment and never across one, so the
`section` and `page` on a chunk are always true rather than approximately true.
"""
import csv
import io
import re
from dataclasses import dataclass

PDF_MAGIC = b"%PDF"
ZIP_MAGIC = b"PK\x03\x04"  # docx and xlsx are both zip containers

_MD_HEADING = re.compile(r"^(#{1,6})\s+(.*)$", re.M)


class UnsupportedDocument(Exception):
    """Raised when nothing here can read the bytes.

    Spring turns this into DocumentStatus.failed with a reason. It must not become a
    document sitting at `ready` with no chunks — `ready` promises retrievable, and a
    document with no chunks can never be retrieved. It would look fine in the UI and
    silently never appear in an answer.
    """


@dataclass(frozen=True)
class Segment:
    """A run of text that shares one piece of provenance."""

    text: str
    section: str | None = None  # Markdown heading, sheet name, or None
    page: int | None = None     # 1-based, PDFs only


@dataclass(frozen=True)
class ExtractedDocument:
    kind: str                 # pdf | markdown | text | csv | xlsx | docx
    segments: list[Segment]
    full_text: str            # segments joined; char offsets on chunks index into THIS

    @property
    def is_empty(self) -> bool:
        return not self.full_text.strip()


def extract(data: bytes, filename: str) -> ExtractedDocument:
    head = data[:8]
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""

    if head.startswith(PDF_MAGIC):
        return _assemble("pdf", _extract_pdf(data))
    if head.startswith(ZIP_MAGIC):
        # Both are zips; only the extension distinguishes them, which is why the real
        # filename has to travel with the bytes.
        if ext == "docx":
            return _assemble("docx", _extract_docx(data))
        if ext in ("xlsx", "xlsm", "xltx"):
            return _assemble("xlsx", _extract_xlsx(data))
        raise UnsupportedDocument(f"zip container with unsupported extension: {ext or '(none)'}")

    text = data.decode("utf-8", errors="replace")
    if ext == "csv":
        return _assemble("csv", _extract_csv(text))
    if ext in ("md", "markdown"):
        return _assemble("markdown", _extract_markdown(text))
    if ext in ("txt", ""):
        return _assemble("text", [Segment(text=text)])
    raise UnsupportedDocument(f"unsupported file type: {ext}")


def _assemble(kind: str, segments: list[Segment]) -> ExtractedDocument:
    kept = [s for s in segments if s.text.strip()]
    return ExtractedDocument(kind=kind, segments=kept, full_text="\n\n".join(s.text for s in kept))


# --- per-format extraction ---------------------------------------------------

def _extract_pdf(data: bytes) -> list[Segment]:
    """One segment per page, so `page` on a citation is real rather than estimated.

    Measured on the live corpus: all 8 PDFs carry a text layer (a 16-page file yields 5,109
    words), so OCR is not needed today. A PDF that extracts to nothing is almost certainly
    scanned; it becomes an empty document and Spring marks it failed with a reason, which is
    the honest signal and the trigger for adding OCR when one actually appears.
    """
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return [Segment(text=(page.extract_text() or ""), page=i)
            for i, page in enumerate(reader.pages, start=1)]


def _extract_markdown(text: str) -> list[Segment]:
    """Split on headings, which hands us `section` for free — no NLP required.

    CitationDto.section already exists end to end and CitationCard.tsx renders it; today it
    shows the literal string "Placeholder section". These headings are what replaces it.
    """
    matches = list(_MD_HEADING.finditer(text))
    if not matches:
        return [Segment(text=text)]

    segments: list[Segment] = []
    preamble = text[: matches[0].start()]
    if preamble.strip():
        segments.append(Segment(text=preamble))
    for i, m in enumerate(matches):
        end = matches[i + 1].start() if i + 1 < len(matches) else len(text)
        segments.append(Segment(text=text[m.start():end], section=m.group(2).strip()))
    return segments


def _extract_csv(text: str) -> list[Segment]:
    """Rows joined with " | ", header repeated is NOT attempted.

    A CSV is not prose and will not chunk like prose. This keeps it readable and citable
    rather than pretending otherwise; whether retrieval over tables is any good is a Day 13
    question, and the honest thing is to find out rather than to reject a third of the
    corpus in advance.
    """
    rows = [" | ".join(c.strip() for c in row) for row in csv.reader(io.StringIO(text)) if any(row)]
    return [Segment(text="\n".join(rows))]


def _extract_xlsx(data: bytes) -> list[Segment]:
    """One segment per sheet; the sheet name becomes `section`.

    Loaded from BytesIO, never a path: openpyxl selects its parser from the file extension,
    and stored files are `uploads/<uuid>` with none — which is the entire reason every
    spreadsheet in the corpus appeared unreadable.
    """
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    segments: list[Segment] = []
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                rows.append(" | ".join(cells))
        segments.append(Segment(text="\n".join(rows), section=ws.title))
    return segments


def _extract_docx(data: bytes) -> list[Segment]:
    """Paragraphs grouped under their most recent Heading style.

    No DOCX exists in the corpus. It is supported because DocumentService accepts it, and
    the first upload should produce chunks rather than a mysterious empty document.
    """
    import docx

    document = docx.Document(io.BytesIO(data))
    segments: list[Segment] = []
    section: str | None = None
    buffer: list[str] = []

    def flush() -> None:
        if buffer:
            segments.append(Segment(text="\n".join(buffer), section=section))
            buffer.clear()

    for para in document.paragraphs:
        if para.style is not None and str(para.style.name).startswith("Heading"):
            flush()
            section = para.text.strip() or section
            buffer.append(para.text)
        elif para.text.strip():
            buffer.append(para.text)
    flush()
    return segments
